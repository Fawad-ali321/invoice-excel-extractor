"""Extract line items from invoice PDFs and verify totals.

Reads every PDF in samples/, pulls vendor/date/line-items/totals with
pdfplumber + regex, then recomputes every invoice total from its line items
plus the stated tax, and writes an Excel
workbook with two sheets: Extracted_Data and Verification.
"""
import os
import re

import pdfplumber
import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
SAMPLES_DIR = os.path.join(BASE, "samples")
OUTPUT_XLSX = os.path.join(BASE, "output", "invoices_extracted.xlsx")

ITEM_RE = re.compile(r"^(.+?)\s+(\d+)\s+([\d,]+\.\d{2})\s+([\d,]+\.\d{2})$")
SUBTOTAL_RE = re.compile(r"Subtotal:\s+([\d,]+\.\d{2})")
TAX_RE = re.compile(r"Tax.*?:\s+([\d,]+\.\d{2})")
TOTAL_RE = re.compile(r"TOTAL:\s+([\d,]+\.\d{2})")
INV_RE = re.compile(r"INVOICE\s+#(\S+)")
DATE_RE = re.compile(r"Date:\s+(\S+)")


def parse_invoice(pdf_path):
    """Return (line-item rows, verification row) for one invoice PDF."""
    with pdfplumber.open(pdf_path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

    vendor = lines[0]
    inv_no = INV_RE.search(text).group(1)
    date = DATE_RE.search(text).group(1)
    stated_subtotal = float(SUBTOTAL_RE.search(text).group(1).replace(",", ""))
    stated_tax = float(TAX_RE.search(text).group(1).replace(",", ""))
    stated_total = float(TOTAL_RE.search(text).group(1).replace(",", ""))

    items = []
    for ln in lines:
        m = ITEM_RE.match(ln)
        if m and not any(k in ln for k in ("Subtotal", "Tax", "TOTAL")):
            desc, qty, price, amount = m.groups()
            items.append({
                "invoice_no": inv_no,
                "date": date,
                "vendor": vendor,
                "description": desc.strip(),
                "qty": int(qty),
                "unit_price": float(price.replace(",", "")),
                "amount": float(amount.replace(",", "")),
            })

    computed_subtotal = round(sum(i["amount"] for i in items), 2)
    computed_total = round(computed_subtotal + stated_tax, 2)
    status = "MATCH" if abs(computed_total - stated_total) < 0.01 else "MISMATCH"
    verification = {
        "invoice_no": inv_no,
        "vendor": vendor,
        "date": date,
        "stated_total": stated_total,
        "computed_total": computed_total,
        "difference": round(stated_total - computed_total, 2),
        "status": status,
    }
    return items, verification


def main():
    all_items, verifications = [], []
    for fname in sorted(os.listdir(SAMPLES_DIR)):
        if fname.lower().endswith(".pdf"):
            items, ver = parse_invoice(os.path.join(SAMPLES_DIR, fname))
            all_items.extend(items)
            verifications.append(ver)

    os.makedirs(os.path.dirname(OUTPUT_XLSX), exist_ok=True)
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(all_items).to_excel(writer, sheet_name="Extracted_Data", index=False)
        pd.DataFrame(verifications).to_excel(writer, sheet_name="Verification", index=False)

        # Colour the status column: green = verified, red = mismatch
        ws = writer.sheets["Verification"]
        headers = list(verifications[0].keys())
        status_col = headers.index("status") + 1
        green = PatternFill("solid", fgColor="C6EFCE")
        red = PatternFill("solid", fgColor="FFC7CE")
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=status_col)
            cell.fill = green if cell.value == "MATCH" else red
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    bad = sum(1 for v in verifications if v["status"] == "MISMATCH")
    print(f"processed {len(verifications)} invoices -> {OUTPUT_XLSX}")
    print(f"verified: {len(verifications) - bad} MATCH, {bad} MISMATCH")


if __name__ == "__main__":
    main()
